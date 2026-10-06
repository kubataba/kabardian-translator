"""Documents in and out: .txt / .md / .docx are read as paragraphs; the translation is written back as .txt or
.docx — the translation alone, or original and translation sentence by sentence with both languages named."""
from __future__ import annotations

import io
from pathlib import Path

READABLE = (".txt", ".md", ".docx")


def read(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        return "\n\n".join(p.text for p in d.paragraphs if p.text.strip())
    if ext not in (".txt", ".md", ""):
        raise ValueError(f"unsupported file type {ext!r}: use .txt, .md or .docx")
    for enc in ("utf-8-sig", "utf-16", "cp1251", "latin-1"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text.replace("\r\n", "\n")
        except UnicodeDecodeError:
            continue
    raise ValueError("cannot decode the file")


def write(translation: str, fmt: str = "txt", title: str | None = None) -> bytes:
    """The translation alone, paragraph for paragraph. fmt: txt | docx."""
    if fmt == "docx":
        import docx
        d = docx.Document()
        if title:
            d.add_heading(title, level=1)
        for p in translation.split("\n"):
            if p.strip():
                d.add_paragraph(p)
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()
    return (translation.rstrip() + "\n").encode("utf-8")


def write_bilingual(pairs: list, src_name: str, tgt_name: str, src: str, tgt: str, fmt: str = "txt",
                    title: str | None = None) -> bytes:
    """`pairs`: one list of (original, translation) per paragraph. The header names both languages; every sentence
    is followed by its translation, each line marked with its language code (.txt) or in its column (.docx)."""
    head = f"{src_name} ({src}) → {tgt_name} ({tgt})"
    if fmt == "docx":
        import docx
        d = docx.Document()
        if title:
            d.add_heading(title, level=1)
        d.add_paragraph(head)
        table = d.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = f"{src_name} ({src})", f"{tgt_name} ({tgt})"
        for cell in table.rows[0].cells:
            cell.paragraphs[0].runs[0].bold = True
        for para in pairs:
            for a, b in para:
                row = table.add_row().cells
                row[0].text, row[1].text = a, b
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()
    lines = ([title, ""] if title else []) + [head, ""]
    for para in pairs:
        for a, b in para:
            lines += [f"[{src}] {a}", f"[{tgt}] {b}", ""]
        lines.append("")
    return ("\n".join(lines).rstrip() + "\n").encode("utf-8")
